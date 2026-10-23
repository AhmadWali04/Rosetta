"""Isolated Postgres acceptance smoke for F1.2/F2.1/F2.2/F2.4–F2.6.

Creates and drops its own temporary local database. Uses real migrations,
DOCX parsing, hybrid memory retrieval, checkpoints, graph and HTTP routes;
only LLM responses and embeddings are fake. Never seeds an existing database.
Run with ``make smoke-hr``. The local database role needs CREATEDB privilege.
"""

from __future__ import annotations

import asyncio
import io
import os
import tempfile
import uuid
from contextlib import ExitStack
from pathlib import Path
from unittest.mock import patch

import psycopg
from alembic import command
from alembic.config import Config
from docx import Document
from psycopg import sql
from sqlalchemy.engine import make_url


def _exercise() -> None:
    # Imports are deliberately delayed until the disposable DB and offline
    # environment are selected; app.db.session constructs its engine at import.
    from fastapi.testclient import TestClient
    from sqlalchemy import select

    from app.db.models import Application, JdRequirement, Job, Profile, ResumeChunk, SkillTaxonomy
    from app.db.session import SessionLocal
    from app.graphs.intake.graph import ensure_checkpointer_schema, run_intake_graph
    from app.graphs.intake.hr_explainer import HrExplanationResult, RequirementExplanation
    from app.graphs.intake.requirement_matcher import RequirementJudgment
    from app.graphs.intake.skill_extractor import SkillExtractionResult
    from app.graphs.intake.skill_normalizer import (
        SkillNormalizationItem,
        SkillNormalizationResult,
    )
    from app.graphs.state import ExtractedSkill
    from app.llm.fake import FakeLLMProvider
    from app.main import app
    from app.rag import namespaces
    from app.rag.embeddings import FakeEmbedder
    from app.rag.indexer import audit_namespace
    from app.rag.vector_store import VectorRecord, get_vector_store
    from app.services.injection_detection import InjectionCheckResult
    from app.services.vector_health import audit_vectors

    job_id, requirement_id, taxonomy_id = "smoke-job", "smoke-python", "smoke-python"
    with SessionLocal() as session:
        session.add(Job(id=job_id, title="Smoke role", raw_jd="Production Python services"))
        session.add(
            JdRequirement(
                id=requirement_id,
                job_id=job_id,
                text="Python services",
                kind="must_have",
                weight=1,
                pinecone_id=namespaces.requirement_record_id(requirement_id),
            )
        )
        session.add(
            SkillTaxonomy(
                id=taxonomy_id,
                canonical="Python",
                alt_labels=[],
                description="Python programming language",
                source="onet",
                pinecone_id=namespaces.skill_record_id(taxonomy_id),
            )
        )
        session.commit()

    async def index_taxonomy() -> None:
        embedder = FakeEmbedder()
        [dense] = await embedder.embed_dense(["Python"])
        [sparse] = await embedder.embed_sparse(["Python"])
        record_id = namespaces.skill_record_id(taxonomy_id)
        namespace = namespaces.taxonomy_namespace("dev")
        metadata = {"schema_version": 1, "embed_model": "smoke-fake"}
        await get_vector_store("dense").upsert(
            namespace, [VectorRecord(id=record_id, values=dense, metadata=metadata)]
        )
        await get_vector_store("sparse").upsert(
            namespace,
            [
                VectorRecord(
                    id=record_id,
                    values=[],
                    sparse_indices=sparse.indices,
                    sparse_values=sparse.values,
                    metadata=metadata,
                )
            ],
        )

    asyncio.run(index_taxonomy())
    ensure_checkpointer_schema()
    headers = {"X-Role": "hr"}
    application_ids = []
    with TestClient(app) as client:
        for index in range(2):
            evidence = f"Built Python services for production project {index}."
            document = Document()
            document.add_heading("Experience", level=1)
            document.add_paragraph(evidence)
            stream = io.BytesIO()
            document.save(stream)
            response = client.post(
                "/applications",
                data={"job_id": job_id, "name": f"Smoke {index}", "email": "fake@example.com"},
                files={
                    "resume": (
                        "synthetic.docx",
                        stream.getvalue(),
                        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    )
                },
            )
            assert response.status_code == 201, "multipart application failed"
            application_id = response.json()["id"]
            application_ids.append(application_id)
            with SessionLocal() as session:
                application = session.get(Application, application_id)
                assert application is not None
                resume_path = application.resume_path
            fake = FakeLLMProvider(
                responses={
                    "InjectionCheckResult": InjectionCheckResult(),
                    "SkillExtractionResult": SkillExtractionResult(
                        skills=[
                            ExtractedSkill(
                                skill="Python",
                                evidence_span=evidence,
                                context="Experience",
                                recency="current",
                                depth_signal="production",
                            )
                        ]
                    ),
                    "SkillNormalizationResult": SkillNormalizationResult(
                        items=[
                            SkillNormalizationItem(
                                skill_index=0, canonical="Python", taxonomy_id=taxonomy_id
                            )
                        ]
                    ),
                    "RequirementJudgment": RequirementJudgment(
                        status="met",
                        cited_chunk_ids=[namespaces.chunk_record_id(application_id, 0)],
                        rationale="The resume describes building Python services.",
                    ),
                    "HrExplanationResult": HrExplanationResult(
                        per_requirement=[
                            RequirementExplanation(
                                requirement_id=requirement_id,
                                explanation="Python services are demonstrated in project work.",
                            )
                        ],
                        overall="",
                    ),
                }
            )
            with ExitStack() as patches:
                # Patch every imported provider accessor, plus the factory cache
                # as a fail-safe for future nodes importing its original accessor.
                patches.enter_context(patch("app.llm.factory._INSTANCES", {"openai": fake}))
                for module in (
                    "app.llm",
                    "app.services.injection_detection",
                    "app.graphs.intake.skill_extractor",
                    "app.graphs.intake.skill_normalizer",
                    "app.graphs.intake.requirement_matcher",
                    "app.graphs.intake.hr_explainer",
                ):
                    patches.enter_context(patch(f"{module}.get_llm_provider", return_value=fake))
                result = run_intake_graph(application_id, job_id, resume_path)
            assert result.vector_status == "indexed", "memory index did not complete"
            assert result.score_components is not None, "score components were not produced"
            with SessionLocal() as session:
                profile = session.get(Profile, application_id)
                assert profile is not None and profile.score_components is not None
            detail = client.get(f"/hr/applications/{application_id}", headers=headers)
            assert detail.status_code == 200
            body = detail.json()
            chunk_ids = {chunk["id"] for chunk in body["chunks"]}
            assert body["requirements"][0]["status"] == "met"
            assert set(body["requirements"][0]["chunk_ids"]) <= chunk_ids
            assert body["requirements"][0]["chunk_ids"], "missing requirement citations"
            assert body["skills"] and body["skills"][0]["canonical"] == "Python"
            label = client.patch(
                f"/applications/{application_id}/status",
                headers=headers,
                json={"status": "pending", "reason": "Offline acceptance review"},
            )
            assert label.status_code == 200

        rows = client.get("/hr/applications", headers=headers, params={"job_id": job_id})
        assert rows.status_code == 200 and len(rows.json()) == 2
        assert all(row["status"] == "pending" for row in rows.json()), "intake undid HR status"
        preview = client.post(
            f"/pools/{job_id}/promote/preview", headers=headers, json={"percent": 50}
        )
        assert preview.status_code == 200 and len(preview.json()["application_ids"]) == 1
        confirmed_ids = preview.json()["application_ids"]
        confirmed = client.post(
            f"/pools/{job_id}/promote",
            headers=headers,
            json={"percent": 50, "application_ids": confirmed_ids},
        )
        assert confirmed.status_code == 200
        for application_id in application_ids:
            detail = client.get(f"/hr/applications/{application_id}", headers=headers).json()
            expected = "valid" if application_id in confirmed_ids else "pending"
            assert detail["status"] == expected
            actions = {event["action"] for event in detail["events"]}
            assert {"intake_complete", "label_changed"} <= actions
            if application_id in confirmed_ids:
                assert "pool_promoted" in actions
        assert client.get("/hr/applications").status_code == 403
    with SessionLocal() as session:
        expected_chunks = list(
            session.scalars(
                select(ResumeChunk.pinecone_id).where(
                    ResumeChunk.application_id.in_(application_ids)
                )
            ).all()
        )

    async def verify_vector_parity() -> None:
        resumes = await audit_namespace(
            namespaces.resumes_namespace("dev", job_id), expected_chunks
        )
        assert not any(
            resumes[key] for key in ("missing", "orphaned", "missing_dense", "missing_sparse")
        ), "resume indexes differ from Postgres chunks"
        summary_ids = [namespaces.candidate_summary_record_id(item) for item in application_ids]
        candidates = await audit_namespace(
            namespaces.candidates_namespace("dev", job_id), summary_ids, expected_sparse_ids=[]
        )
        assert candidates["parity"], "candidate indexes differ from their explicit storage policy"
        # Current intake deliberately stores candidate summaries in dense only.
        # Verify that policy explicitly rather than weakening dual-index checks
        # for namespaces which do require both representations.
        sparse_summaries = await get_vector_store("sparse").list_ids(
            namespaces.candidates_namespace("dev", job_id)
        )
        assert not sparse_summaries, "unexpected sparse candidate summaries"
        with SessionLocal() as session:
            for namespace in (
                namespaces.resumes_namespace("dev", job_id),
                namespaces.candidates_namespace("dev", job_id),
            ):
                reports = await audit_vectors(session, namespace)
                assert len(reports) == 1 and reports[0]["parity"], (
                    "Postgres-derived vector audit failed"
                )

    asyncio.run(verify_vector_parity())
    print("PASS: migrations, DOCX upload, real checkpointed intake, citations, pool scoring,")
    print("      HR labels, promotion confirmation and audit (2 synthetic applications).")
    print("PASS: resume vector parity in both indexes; candidate summary parity in dense index.")
    if os.environ.get("SKILLBRIDGE_SMOKE_COPILOT") == "1":
        from scripts.smoke_copilot import exercise

        exercise(job_id, application_ids)


def main() -> None:
    # Keep all external inference/tracing disabled even when the caller loads .env.
    os.environ.update(
        VECTOR_BACKEND="memory",
        PINECONE_ENV="dev",
        LLM_PROVIDER="openai",
        RERANK_ENABLED="false",
        LANGFUSE_PUBLIC_KEY="",
        LANGFUSE_SECRET_KEY="",
        LANGCHAIN_TRACING_V2="false",
        LANGSMITH_TRACING="false",
    )
    from app.config import get_settings

    get_settings.cache_clear()
    original_url = make_url(get_settings().database_url)
    if original_url.get_backend_name() != "postgresql" or original_url.host not in {
        "localhost",
        "127.0.0.1",
        "::1",
    }:
        raise RuntimeError("smoke-hr requires a local PostgreSQL DATABASE_URL")
    name = f"skillbridge_smoke_{uuid.uuid4().hex}"
    admin_url = original_url.set(drivername="postgresql", database="postgres")
    smoke_url = original_url.set(database=name)
    created = False
    with psycopg.connect(admin_url.render_as_string(hide_password=False), autocommit=True) as admin:
        try:
            admin.execute(sql.SQL("CREATE DATABASE {}").format(sql.Identifier(name)))
            created = True
            os.environ["DATABASE_URL"] = smoke_url.render_as_string(hide_password=False)
            get_settings.cache_clear()
            api_root = Path(__file__).resolve().parents[1]
            alembic = Config(str(api_root / "alembic.ini"))
            alembic.set_main_option("script_location", str(api_root / "app/db/migrations"))
            command.upgrade(alembic, "head")
            command.check(alembic)
            with tempfile.TemporaryDirectory(prefix="skillbridge-smoke-") as upload_dir:
                os.environ["RESUME_UPLOAD_DIR"] = upload_dir
                get_settings.cache_clear()
                # TestClient has its own in-process transport; any actual HTTP
                # request from an LLM/vector SDK must fail rather than leave host.
                with ExitStack() as guards:
                    for transport in (
                        "httpx.HTTPTransport.handle_request",
                        "httpx.AsyncHTTPTransport.handle_async_request",
                        "urllib3.connectionpool.HTTPConnectionPool.urlopen",
                    ):
                        guards.enter_context(
                            patch(transport, side_effect=RuntimeError("external HTTP disabled"))
                        )
                    _exercise()
        finally:
            if created:
                # Only this freshly created, UUID-named database may be dropped.
                import sys

                db_module = sys.modules.get("app.db.session")
                if db_module is not None:
                    db_module.engine.dispose()
                admin.execute(sql.SQL("DROP DATABASE {} WITH (FORCE)").format(sql.Identifier(name)))
                print("Temporary smoke database removed; existing databases untouched.")


if __name__ == "__main__":
    main()
