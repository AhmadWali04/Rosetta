"""PDF/DOCX resume extraction (PRD F1.2: resume upload, PDF/DOCX, <= 5MB;
PRD §7.1 "Ingest & Parse": "PDF/DOCX -> text + section segmentation").

Uses `pdfplumber` for PDF (exposes per-character style metadata —
`non_stroking_color`, font `size` — which is what makes hidden-text
detection possible; `pypdf` alone doesn't expose per-char style, so it's
not used here) and `python-docx` for DOCX (exposes `run.font.color.rgb`
and `run.font.hidden` directly).

Hidden-text detection lives HERE, not in the sanitize node
(`app/graphs/intake/sanitize.py`, another agent's file). `extract_pdf`/
`extract_docx` return a `RawDoc` whose `hidden_spans` already identifies
exactly which spans of `text` were hidden (white-on-white / tiny-font for
PDF, `font.hidden`/white color for DOCX) and WHY (the offending text is
preserved, never silently dropped, for an audit trail). This keeps the one
piece of logic that understands "hidden" in one well-tested place instead
of splitting it awkwardly between this module and the sanitize node, which
only needs to consume `hidden_spans` and decide what to do about them
(strip + flag).

`segment_sections` is a pragmatic v1 heuristic — see its docstring.
"""

from __future__ import annotations

import re

from pydantic import BaseModel, Field

from app.graphs.state import Section

# --- Hidden-text thresholds -------------------------------------------------

# A fill color is "white/near-white" if every RGB channel (0.0-1.0 scale,
# pdfplumber's `non_stroking_color` convention) is at or above this.
_WHITE_THRESHOLD = 0.92
# A grayscale fill (single-channel `non_stroking_color`) counts as
# white/near-white at or above this.
_GRAY_WHITE_THRESHOLD = 0.95
# A CMYK fill counts as white/near-white when every channel is at or below
# this (CMYK: 0 ink on every channel = white).
_CMYK_WHITE_THRESHOLD = 0.08
# Font size (points) below which text is considered "hidden" even if not
# white — effectively invisible at normal zoom.
_TINY_FONT_PT = 1.0
# DOCX: an RGB channel (0-255) at or above this, on all three channels,
# counts as white/near-white.
_DOCX_WHITE_CHANNEL = 0xF0


class StyledSpan(BaseModel):
    """One span of text with a style-derived `is_hidden` verdict already
    computed. `char_start`/`char_end` are offsets into the `RawDoc.text`
    this span came from. The offending `text` is always preserved (never
    dropped without a record) so sanitize/audit logic downstream can show
    what was found and removed.
    """

    text: str
    char_start: int
    char_end: int
    is_hidden: bool


class RawDoc(BaseModel):
    """Result of extracting one resume file."""

    text: str
    hidden_spans: list[StyledSpan] = Field(default_factory=list)


def _pdf_color_is_white(color: object) -> bool:
    """`non_stroking_color` from pdfplumber: `None`, a single float
    (grayscale), or a tuple of 3 (RGB) or 4 (CMYK) floats in 0.0-1.0."""
    if color is None:
        return False
    if isinstance(color, int | float):
        return color >= _GRAY_WHITE_THRESHOLD
    if isinstance(color, tuple | list):
        if len(color) == 1:
            return color[0] >= _GRAY_WHITE_THRESHOLD
        if len(color) == 3:
            return all(c >= _WHITE_THRESHOLD for c in color)
        if len(color) == 4:
            return all(c <= _CMYK_WHITE_THRESHOLD for c in color)
    return False


def extract_pdf(path: str) -> RawDoc:
    """Extract text and hidden spans from a PDF via `pdfplumber`.

    Builds `text` from each page's `extract_text()` (joined with a blank
    line between pages) — this deliberately includes hidden text, the same
    way a naive "just extract the text" pipeline would, because that's
    exactly the vulnerability: hidden-from-a-human text still ends up in
    the extracted string unless something flags it. Separately, it scans
    each page's `chars` for runs of hidden-looking characters (white/near-
    white fill color, or a tiny font size) and locates each run inside the
    assembled `text` via substring search to compute `char_start`/
    `char_end`. The substring-search step is a v1 heuristic: if an exact
    match can't be found (e.g. due to whitespace differences between
    `extract_text()`'s layout reconstruction and the raw char stream), the
    span is still recorded with `char_start=char_end=-1` rather than
    silently dropped.
    """
    import pdfplumber

    page_texts: list[str] = []
    hidden_runs: list[str] = []

    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            page_texts.append(page.extract_text() or "")

            current_run: list[str] = []
            for char in page.chars:
                is_hidden = (
                    _pdf_color_is_white(char.get("non_stroking_color"))
                    or char.get("size", _TINY_FONT_PT + 1) < _TINY_FONT_PT
                )
                if is_hidden:
                    current_run.append(char.get("text", ""))
                elif current_run:
                    hidden_runs.append("".join(current_run))
                    current_run = []
            if current_run:
                hidden_runs.append("".join(current_run))

    text = "\n\n".join(page_texts)
    hidden_spans = _locate_hidden_runs(text, hidden_runs)
    return RawDoc(text=text, hidden_spans=hidden_spans)


def _locate_hidden_runs(text: str, hidden_runs: list[str]) -> list[StyledSpan]:
    spans: list[StyledSpan] = []
    cursor = 0
    for run in hidden_runs:
        stripped = run.strip()
        if not stripped:
            continue
        idx = text.find(stripped, cursor)
        if idx == -1:
            idx = text.find(stripped)
        if idx == -1:
            # Couldn't locate it exactly (layout-reconstruction whitespace
            # differences) — record it anyway for the audit trail, with
            # offsets marked unknown rather than dropping it silently.
            spans.append(StyledSpan(text=stripped, char_start=-1, char_end=-1, is_hidden=True))
            continue
        end = idx + len(stripped)
        spans.append(StyledSpan(text=stripped, char_start=idx, char_end=end, is_hidden=True))
        cursor = end
    return spans


def _docx_run_is_hidden(run: object) -> bool:
    font = run.font  # type: ignore[attr-defined]
    if font.hidden:
        return True
    try:
        color = font.color
        rgb = color.rgb if color is not None and color.type is not None else None
    except (AttributeError, KeyError):
        rgb = None
    if rgb is None:
        return False
    # python-docx's RGBColor is a 3-byte sequence; index access gives ints.
    return all(channel >= _DOCX_WHITE_CHANNEL for channel in (rgb[0], rgb[1], rgb[2]))


def extract_docx(path: str) -> RawDoc:
    """Extract text and hidden spans from a DOCX via `python-docx`.

    Walks paragraphs and their runs directly (rather than a single
    `document.paragraphs` text join) so each run's hidden/visible style can
    be checked individually and its exact `char_start`/`char_end` in the
    assembled `text` is known precisely — no substring search needed,
    unlike the PDF path.
    """
    from docx import Document

    document = Document(path)
    text_parts: list[str] = []
    hidden_spans: list[StyledSpan] = []
    cursor = 0

    for paragraph in document.paragraphs:
        for run in paragraph.runs:
            run_text = run.text
            if not run_text:
                continue
            start = cursor
            end = start + len(run_text)
            if _docx_run_is_hidden(run):
                hidden_spans.append(
                    StyledSpan(text=run_text, char_start=start, char_end=end, is_hidden=True)
                )
            text_parts.append(run_text)
            cursor = end
        text_parts.append("\n")
        cursor += 1

    text = "".join(text_parts).rstrip("\n")
    return RawDoc(text=text, hidden_spans=hidden_spans)


# --- Section segmentation ---------------------------------------------------

# PRD §7.1 "Ingest & Parse: ... section segmentation (Experience, Projects,
# Skills, Education)." This is a pragmatic v1 heuristic, NOT real NLP
# segmentation: a line is a heading candidate if it's short and
# ALL-CAPS-or-Title-Case AND matches one of a fixed set of common resume
# heading words (case-insensitively, punctuation-stripped). Resumes that
# use unlisted heading words, multi-line headings, or unusual formatting
# will fall back to the single "body" section below. Good enough to
# demonstrate the pipeline; not a substitute for real document structure
# parsing.
_HEADING_MAX_LEN = 40
_HEADING_KEYWORDS = {
    "experience",
    "work experience",
    "professional experience",
    "projects",
    "skills",
    "technical skills",
    "education",
    "summary",
    "objective",
    "certifications",
    "publications",
}
_PUNCTUATION_RE = re.compile(r"[^\w\s]")


def _normalize_heading(line: str) -> str:
    return _PUNCTUATION_RE.sub("", line).strip().lower()


def _is_heading_candidate(line: str) -> bool:
    stripped = line.strip()
    if not stripped or len(stripped) >= _HEADING_MAX_LEN:
        return False
    if _normalize_heading(stripped) not in _HEADING_KEYWORDS:
        return False
    has_letters = any(c.isalpha() for c in stripped)
    is_all_caps = has_letters and stripped.upper() == stripped
    is_title_case = stripped == stripped.title()
    return is_all_caps or is_title_case


def segment_sections(text: str) -> list[Section]:
    """Heuristic section detection (see module-level note above).

    Scans line-by-line for heading candidates and splits `text` into one
    `Section` per heading, each spanning from that heading's line start to
    the next heading's line start (or the end of the document), with
    correct absolute offsets into `text`. Any text before the first
    detected heading becomes a leading `Section(name="body", ...)`.

    If zero headings are detected, the entire document becomes one
    `Section(name="body", text=text, char_start=0, char_end=len(text))` —
    this function never returns an empty list.
    """
    headings: list[tuple[int, str]] = []  # (char offset of line start, heading text)
    cursor = 0
    for line in text.splitlines(keepends=True):
        content = line.splitlines()[0] if line.splitlines() else ""
        if _is_heading_candidate(content):
            headings.append((cursor, content.strip()))
        cursor += len(line)

    if not headings:
        return [Section(name="body", text=text, char_start=0, char_end=len(text))]

    sections: list[Section] = []

    first_heading_start = headings[0][0]
    if first_heading_start > 0 and text[:first_heading_start].strip():
        sections.append(
            Section(
                name="body",
                text=text[:first_heading_start],
                char_start=0,
                char_end=first_heading_start,
            )
        )

    boundaries = [h[0] for h in headings] + [len(text)]
    for i, (start, name) in enumerate(headings):
        end = boundaries[i + 1]
        sections.append(Section(name=name, text=text[start:end], char_start=start, char_end=end))

    return sections
